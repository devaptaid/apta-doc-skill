#!/usr/bin/env python3
"""Export Aptaworks proposal HTML to editable DOCX after explicit approval."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from urllib.parse import unquote, urlparse


NAVY = "17234E"
DEEP_NAVY = "0B1533"
ORANGE = "EF7636"
MUTED = "5E6B82"
RULE = "C9D2DF"
WARM = "F7F5EE"


def _dependencies():
    try:
        from bs4 import BeautifulSoup
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Missing DOCX dependency. See references/conversion.md.") from exc
    return BeautifulSoup, Document


def _shade(element, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    properties = element._tc.get_or_add_tcPr() if hasattr(element, "_tc") else element._p.get_or_add_pPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    properties.append(node)


def _set_font(target, name: str) -> None:
    """Set Word font names explicitly so theme fonts cannot override them."""
    from docx.oxml.ns import qn
    target.font.name = name
    r_pr = target._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), name)


def _set_repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def _set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    """Keep tblW, tblInd, tblGrid, and every tcW consistent in DXA."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def _set_cover_table_geometry(table) -> None:
    """Match the full-bleed 100% cover table used by proposal-simple.docx."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    column = OxmlElement("w:gridCol")
    column.set(qn("w:w"), "11906")
    grid.append(column)
    tc_w = table.cell(0, 0)._tc.get_or_add_tcPr().find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        table.cell(0, 0)._tc.get_or_add_tcPr().append(tc_w)
    tc_w.set(qn("w:w"), "5000")
    tc_w.set(qn("w:type"), "pct")


def _add_page_field(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _append_field(run, instruction: str, result: str | None = None, *, end: bool = True) -> None:
    """Append a Word field with an optional cached result to an existing run."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, instr, separate])
    if result is not None:
        text = OxmlElement("w:t")
        text.text = str(result)
        run._r.append(text)
    if end:
        _append_field_end(run)


def _append_field_end(run) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _add_internal_hyperlink(paragraph, anchor: str, runs: list) -> None:
    """Wrap TOC runs in a same-document hyperlink to a bookmark."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", anchor):
        raise ValueError(f"Internal hyperlink anchor must be a bookmark name, not a file path: {anchor}")
    hyperlink = OxmlElement("w:hyperlink")
    # Intentionally set only w:anchor. A same-document hyperlink must not
    # receive r:id, which would create an external/file relationship.
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    for run in runs:
        xml_run = run._r
        xml_run.getparent().remove(xml_run)
        hyperlink.append(xml_run)
    paragraph._p.append(hyperlink)


def _add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    """Wrap a heading paragraph in a bookmark for TOC PAGEREF links."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p = paragraph._p
    insert_at = 1 if p.pPr is not None else 0
    p.insert(insert_at, start)
    p.append(end)


def _disable_update_fields(doc) -> None:
    """Prevent Word's open-time field-security prompt.

    The TOC and PAGEREF fields retain cached results and can be refreshed
    manually in Word. Setting ``w:updateFields`` to true makes Word ask for
    permission to update fields as the document opens, even for internal
    references.
    """
    from docx.oxml.ns import qn

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is not None:
        settings.remove(update)


def _extract_toc_items(soup) -> list[dict[str, str | int]]:
    items = []
    for entry in soup.select(".toc-entry"):
        items.append({
            "level": int(entry.get("data-toc-level", "1")),
            "number": entry.get("data-toc-number", ""),
            "text": entry.get("data-toc-text", ""),
            "display": entry.get("data-toc-display", entry.get_text(" ", strip=True)),
            "id": entry.get("data-toc-id", ""),
        })
    return items


def _add_table_of_contents(doc, toc_items: list[dict[str, str | int]], page_map: dict[str, int], title_text: str = "DAFTAR ISI") -> None:
    """Add the sample's editable Word TOC with cached page references."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Cm, Pt, RGBColor

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run(title_text)
    _set_font(run, "Georgia")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    _set_paragraph_bottom_border(title, ORANGE, size=8)

    field_anchor = doc.add_paragraph()
    field_anchor.paragraph_format.space_after = Pt(0)
    anchor_run = field_anchor.add_run()
    _append_field(anchor_run, r'TOC \o "1-3" \h \z \u', end=False)

    last_entry = None
    for index, item in enumerate(toc_items, start=1):
        level = int(item["level"])
        display = str(item["display"])
        toc_paragraph = doc.add_paragraph(style=f"TOC {min(level, 3)}")
        toc_paragraph.paragraph_format.left_indent = Cm(0.4 * max(level - 1, 0))
        toc_paragraph.paragraph_format.space_before = Pt(4 if level > 1 else 8)
        toc_paragraph.paragraph_format.space_after = Pt(3)
        toc_paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        title_run = toc_paragraph.add_run(display)
        _set_font(title_run, "Lato")
        title_run.font.size = Pt(11 if level == 1 else 10)
        title_run.font.bold = level == 1
        title_run.font.color.rgb = RGBColor.from_string(NAVY)
        tab_run = toc_paragraph.add_run("\t")
        page_run = toc_paragraph.add_run()
        _set_font(page_run, "Lato")
        page_run.font.size = Pt(11 if level == 1 else 10)
        page_run.font.color.rgb = RGBColor.from_string(MUTED)
        bookmark = f"_TocApta{index}"
        cached_page = page_map.get(display, "")
        _append_field(page_run, f"PAGEREF {bookmark} \\h", str(cached_page))
        _add_internal_hyperlink(toc_paragraph, bookmark, [title_run, tab_run, page_run])
        last_entry = toc_paragraph

    end_run = (last_entry or field_anchor).add_run()
    _append_field_end(end_run)


def extract_heading_page_numbers(pdf_path: Path, html_path: Path) -> dict[str, int]:
    """Extract body-page locations for the cached DOCX TOC results."""
    try:
        from bs4 import BeautifulSoup
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Missing PDF/DOCX dependency. See references/conversion.md.") from exc

    soup = BeautifulSoup(html_path.read_text(encoding="utf-8"), "html.parser")
    items = _extract_toc_items(soup)
    if not items:
        return {}

    def normalize(value: str) -> str:
        return re.sub(r"\s+", " ", value or "").strip().lower()

    page_texts = [normalize(page.extract_text() or "") for page in PdfReader(str(pdf_path)).pages]
    first_heading = normalize(str(items[0]["display"]))
    occurrences = [index for index, text in enumerate(page_texts) if first_heading in text]
    body_start = occurrences[1] if len(occurrences) > 1 else (occurrences[0] + 1 if occurrences else 1)

    page_map: dict[str, int] = {}
    for item in items:
        display = str(item["display"])
        target = normalize(display)
        for page_index in range(body_start, len(page_texts)):
            if target in page_texts[page_index]:
                page_map[display] = page_index + 1
                break
    return page_map


def _set_paragraph_bottom_border(paragraph, color: str, size: int = 6) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def _local_image(src: str, html_dir: Path) -> Path | None:
    parsed = urlparse(src)
    if parsed.scheme == "file":
        raw = unquote(parsed.path)
        if re.match(r"^/[A-Za-z]:/", raw):
            raw = raw[1:]
        path = Path(raw)
    elif parsed.scheme:
        return None
    else:
        path = (html_dir / unquote(src)).resolve()
    return path if path.exists() else None


def _format_inline_run(run, element, inherited_elements=None) -> None:
    """Preserve basic inline emphasis and explicit HTML text colors in DOCX."""
    from docx.shared import RGBColor

    descendants = [*(inherited_elements or []), element, *element.find_all(True)]
    run.bold = any(tag.name in {"strong", "b"} for tag in descendants)
    run.italic = any(tag.name in {"em", "i"} for tag in descendants)
    if any("c62828" in (tag.get("style", "").lower().replace(" ", "")) for tag in descendants):
        run.font.color.rgb = RGBColor.from_string("C62828")
    if element.name == "code":
        run.font.name = "Consolas"


def _text_with_runs(paragraph, element) -> None:
    from bs4 import NavigableString, Tag
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for child in element.children:
        if isinstance(child, NavigableString):
            if str(child):
                paragraph.add_run(str(child))
        elif isinstance(child, Tag):
            if child.name == "br":
                paragraph.add_run().add_break()
                continue
            run = paragraph.add_run(child.get_text(" "))
            _format_inline_run(run, child, [element])


def _text_with_runs_excluding_lists(paragraph, element) -> None:
    """Add inline cell content without flattening nested list text into a run."""
    from bs4 import NavigableString, Tag
    for child in element.children:
        if isinstance(child, NavigableString):
            if str(child):
                paragraph.add_run(str(child))
        elif isinstance(child, Tag):
            if child.name in {"ul", "ol"}:
                continue
            if child.name == "br":
                paragraph.add_run().add_break()
                continue
            run = paragraph.add_run(child.get_text(" "))
            _format_inline_run(run, child, [element])


def _set_cell_paragraph_format(paragraph, row_index: int) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.size = Pt(10)
        if row_index == 0:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)


def _populate_table_cell(cell, source_cell, row_index: int) -> None:
    """Populate a DOCX table cell while preserving real HTML lists as Word lists."""
    from bs4 import NavigableString, Tag

    cell.text = ""
    first_paragraph = cell.paragraphs[0]
    first_paragraph.clear()
    paragraph_used = False

    def add_paragraph(style=None):
        nonlocal paragraph_used
        if not paragraph_used:
            paragraph = first_paragraph
            if style:
                paragraph.style = style
            paragraph_used = True
            return paragraph
        return cell.add_paragraph(style=style)

    def add_list(list_element, level=0):
        style = "List Bullet" if list_element.name == "ul" else "List Number"
        for item in list_element.find_all("li", recursive=False):
            paragraph = add_paragraph(style)
            _text_with_runs_excluding_lists(paragraph, item)
            _set_cell_paragraph_format(paragraph, row_index)
            for nested in item.find_all(["ul", "ol"], recursive=False):
                add_list(nested, level + 1)

    for child in source_cell.children:
        if isinstance(child, NavigableString):
            if str(child).strip():
                paragraph = add_paragraph()
                paragraph.add_run(str(child).strip())
                _set_cell_paragraph_format(paragraph, row_index)
        elif isinstance(child, Tag):
            if child.name in {"ul", "ol"}:
                add_list(child)
            elif child.name == "p":
                paragraph = add_paragraph()
                _text_with_runs_excluding_lists(paragraph, child)
                _set_cell_paragraph_format(paragraph, row_index)
            else:
                paragraph = add_paragraph()
                _text_with_runs_excluding_lists(paragraph, child)
                _set_cell_paragraph_format(paragraph, row_index)

    if not paragraph_used:
        _set_cell_paragraph_format(first_paragraph, row_index)


def _table_widths(source_table, column_count: int) -> list[int]:
    """Choose readable DXA widths from the table's semantic column headers."""
    total = 9072
    rows = source_table.find_all("tr")
    header_cells = rows[0].find_all(["th", "td"], recursive=False) if rows else []
    headers = [" ".join(cell.get_text(" ", strip=True).lower().split()) for cell in header_cells]
    key = tuple(headers)

    profiles = {
        ("kebutuhan", "respons solusi"): [2800, 6272],
        ("temuan aptaworks", "implikasi terhadap delivery"): [4200, 4872],
        ("persona", "workload/tool yang direkomendasikan", "peran dalam solusi ckb"): [1500, 4000, 3572],
        ("kelompok", "item data/tabel yang menjadi referensi fin", "status baseline dan catatan"): [1800, 3800, 3472],
        ("halaman/area", "kapabilitas utama"): [2700, 6372],
        ("id", "persyaratan", "cakupan yang diusulkan", "status"): [700, 2200, 4800, 1372],
        ("aktivitas", "mode delivery", "pengaturan yang diusulkan"): [3000, 1700, 4372],
        ("periode", "fase", "aktivitas dan checkpoint", "dependensi utama", "output utama"): [1050, 1800, 3300, 1700, 1222],
        ("fase", "periode", "total md"): [4000, 1500, 3572],
        ("peran", "personel kunci", "tanggung jawab", "alokasi baseline", "lokasi"): [1500, 1900, 3300, 1200, 1172],
        ("id", "deliverable", "isi minimum", "kriteria penerimaan"): [700, 1800, 3000, 3572],
        ("proyek / client", "platform / sumber data", "scope dan domain", "relevansi untuk ckb"): [2200, 1800, 3000, 2072],
        ("item komersial", "isi yang perlu dilengkapi"): [2500, 6572],
        ("risiko", "pengendalian yang diusulkan"): [3000, 6072],
        ("klasifikasi", "contoh perubahan", "perlakuan"): [1600, 4200, 3272],
        ("key", "information"): [2500, 6572],
    }
    if key in profiles:
        widths = profiles[key]
        if len(widths) == column_count and sum(widths) == total:
            return widths

    # Fallback for future tables: estimate relative demand while enforcing
    # readable minimums and a fixed total compatible with table_geometry.py.
    if column_count == 1:
        return [total]
    sample_rows = rows[1:4]
    demands = []
    for index in range(column_count):
        header_len = len(headers[index]) if index < len(headers) else 0
        body_len = 0
        for row in sample_rows:
            cells = row.find_all(["th", "td"], recursive=False)
            if index < len(cells):
                body_len = max(body_len, len(cells[index].get_text(" ", strip=True)))
        demands.append(max(12, min(80, header_len * 1.5, body_len / 3 or 12)))
    minimum = 720
    available = total - minimum * column_count
    demand_total = sum(demands) or column_count
    widths = [minimum + round(available * demand / demand_total) for demand in demands]
    widths[-1] += total - sum(widths)
    return widths


def _style_document(doc) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    styles = doc.styles
    normal = styles["Normal"]
    _set_font(normal, "Lato")
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, before, after in (("Title", 32, 0, 12), ("Heading 1", 24, 10, 10), ("Heading 2", 17, 16, 6), ("Heading 3", 12, 12, 4)):
        style = styles[name]
        _set_font(style, "Georgia" if name != "Heading 3" else "Lato")
        style.font.size = Pt(size)
        style.font.bold = name == "Heading 3"
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for level, name in enumerate(("Heading 1", "Heading 2", "Heading 3")):
        p_pr = styles[name]._element.get_or_add_pPr()
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            p_pr.append(outline)
        outline.set(qn("w:val"), str(level))
    for name, size, bold in (("TOC 1", 11, True), ("TOC 2", 10, False), ("TOC 3", 10, False)):
        if name not in styles:
            toc_style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            toc_style = styles[name]
        _set_font(toc_style, "Lato")
        toc_style.font.size = Pt(size)
        toc_style.font.bold = bold
        toc_style.font.color.rgb = RGBColor.from_string(NAVY)
        toc_style.paragraph_format.space_after = Pt(3)
    if "Apta Code" not in styles:
        code = styles.add_style("Apta Code", WD_STYLE_TYPE.PARAGRAPH)
        _set_font(code, "Consolas")
        code.font.size = Pt(9)
        code.font.color.rgb = RGBColor.from_string(NAVY)
        code.paragraph_format.left_indent = Pt(12)
        code.paragraph_format.right_indent = Pt(12)
        code.paragraph_format.space_before = Pt(6)
        code.paragraph_format.space_after = Pt(8)


def _add_cover(doc, meta: dict[str, str], logo_white: Path, soup, html_dir: Path) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(0)
    section.header_distance = section.footer_distance = Cm(0)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    row = table.rows[0]
    row.height = Cm(29.7)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = row.cells[0]
    cell.width = Cm(21)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _shade(cell, DEEP_NAVY)
    _set_cell_margins(cell, top=0, start=0, bottom=0, end=0)

    def cover_paragraph(paragraph, *, before=0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT):
        paragraph.alignment = align
        paragraph.paragraph_format.left_indent = Cm(2.5)
        paragraph.paragraph_format.right_indent = Cm(2.5)
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        return paragraph

    logo = cell.paragraphs[0]
    cover_paragraph(logo, before=85, after=0)
    logo.add_run().add_picture(str(logo_white), width=Cm(4.6))
    label = cover_paragraph(cell.add_paragraph(), before=70, after=12)
    run = label.add_run("PROPOSAL")
    _set_font(run, "Consolas")
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(ORANGE)
    title = cover_paragraph(cell.add_paragraph(), after=18)
    title.paragraph_format.line_spacing = 1.05
    run = title.add_run(meta.get("title", "Aptaworks Proposal"))
    _set_font(run, "Georgia")
    run.font.size, run.font.bold = Pt(32), True
    run.font.color.rgb = RGBColor(255, 255, 255)
    subtitle = cover_paragraph(cell.add_paragraph(), after=18)
    run = subtitle.add_run(meta.get("subtitle", ""))
    _set_font(run, "Lato")
    run.font.size = Pt(13.5)
    run.font.color.rgb = RGBColor.from_string("C9D2DF")

    prepared = cover_paragraph(cell.add_paragraph(), before=12, after=4)
    run = prepared.add_run("PREPARED FOR")
    _set_font(run, "Consolas")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("96A0B9")

    client_mark = soup.select_one(".client-logo-box img, .client-mark img")
    client_path = _local_image(client_mark.get("src", ""), html_dir) if client_mark else None
    if client_path:
        client = cover_paragraph(cell.add_paragraph(), after=35)
        client.add_run().add_picture(str(client_path), height=Cm(2.45))
    else:
        client_table = cell.add_table(rows=1, cols=1)
        client_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        client_table.autofit = False
        client_cell = client_table.cell(0, 0)
        _shade(client_cell, "FFFFFF")
        _set_cell_margins(client_cell, top=150, start=220, bottom=150, end=220)
        mark = client_cell.paragraphs[0]
        mark.paragraph_format.space_after = Pt(0)
        run = mark.add_run(meta.get("client", "Client"))
        _set_font(run, "Lato")
        run.font.size, run.font.bold = Pt(11), True
        run.font.color.rgb = RGBColor.from_string(NAVY)
        _set_table_geometry(client_table, [2900], indent=1417)
        trailing = cell.paragraphs[-1]
        cover_paragraph(trailing, after=0)

    spacer = cover_paragraph(cell.add_paragraph(), before=92, after=0)
    watermark = cover_paragraph(cell.add_paragraph(), after=40, align=WD_ALIGN_PARAGRAPH.RIGHT)
    run = watermark.add_run("APTA")
    _set_font(run, "Georgia")
    run.font.size, run.font.italic = Pt(90), True
    run.font.color.rgb = RGBColor.from_string("1E294B")
    footer = cover_paragraph(cell.add_paragraph(), before=12, align=WD_ALIGN_PARAGRAPH.RIGHT)
    run = footer.add_run(
        f"{'Author' if meta.get('language') == 'en' else 'Penulis'}: {meta.get('author', 'Aptaworks Team')}\n"
        f"{'Date' if meta.get('language') == 'en' else 'Tanggal'}: {meta.get('date', '')} (v{meta.get('version', '1.0')})"
    )
    _set_font(run, "Consolas")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("96A0B9")
    _set_cover_table_geometry(table)


def _add_content(
    doc,
    soup,
    html_dir: Path,
    meta: dict[str, str],
    logo_navy: Path,
    page_map: dict[str, int] | None = None,
) -> None:
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt, RGBColor
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(3), Cm(3)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2.5)
    section.header_distance, section.footer_distance = Cm(1.2), Cm(1.2)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.clear()
    header.paragraph_format.tab_stops.add_tab_stop(Cm(15.7), WD_TAB_ALIGNMENT.RIGHT)
    header.add_run(f"{meta.get('status', 'Confidential')} - {meta.get('title', '')}")
    header.add_run("\t")
    header.add_run().add_picture(str(logo_navy), width=Cm(3.5))
    header.style = doc.styles["Normal"]
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    header.runs[0].font.bold = True
    _set_paragraph_bottom_border(header, ORANGE)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Aptaworks Team  |  Page ")
    _add_page_field(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    main = soup.select_one("main") or soup.body
    toc_items = _extract_toc_items(soup)
    bookmark_by_id = {
        str(item["id"]): f"_TocApta{index}"
        for index, item in enumerate(toc_items, start=1)
        if item.get("id")
    }
    for element in main.find_all(recursive=False):
        name = element.name
        if name == "div" and "toc-page" in element.get("class", []):
            _add_table_of_contents(doc, toc_items, page_map or {}, "CONTENTS" if meta.get('language') == 'en' else "DAFTAR ISI")
            doc.add_page_break()
        elif name in {"h1", "h2", "h3"}:
            paragraph = doc.add_paragraph(
                element.get_text(" ", strip=True),
                style={"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[name],
            )
            bookmark_name = bookmark_by_id.get(element.get("id", ""))
            if bookmark_name:
                bookmark_id = list(bookmark_by_id.values()).index(bookmark_name) + 1
                _add_bookmark(paragraph, bookmark_name, bookmark_id)
        elif name == "p":
            if "document-end" in element.get("class", []):
                continue
            images = element.find_all("img", recursive=False)
            if images:
                for image in images:
                    path = _local_image(image.get("src", ""), html_dir)
                    if path:
                        paragraph = doc.add_paragraph()
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        paragraph.add_run().add_picture(str(path), width=Cm(14.5))
                text_content = element.get_text(" ", strip=True)
                if text_content:
                    paragraph = doc.add_paragraph()
                    _text_with_runs(paragraph, element)
            else:
                paragraph = doc.add_paragraph()
                _text_with_runs(paragraph, element)
        elif name in {"ul", "ol"}:
            list_style = "List Bullet" if name == "ul" else "List Number"
            items = element.find_all("li", recursive=False)
            for index, item in enumerate(items):
                paragraph = doc.add_paragraph(item.get_text(" ", strip=True), style=list_style)
                paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if len(items) <= 6:
                    paragraph.paragraph_format.keep_with_next = index < len(items) - 1
        elif name == "blockquote":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(.45)
            paragraph.paragraph_format.right_indent = Cm(.25)
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _shade(paragraph, WARM)
            _text_with_runs(paragraph, element)
            for run in paragraph.runs:
                if not run.font.color.rgb:
                    run.font.color.rgb = RGBColor.from_string(NAVY)
        elif name == "pre":
            paragraph = doc.add_paragraph(element.get_text("\n", strip=False), style="Apta Code")
            _shade(paragraph, "F1F5F9")
        elif name == "img":
            path = _local_image(element.get("src", ""), html_dir)
            if path:
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(path), width=Cm(14.5))
        elif name == "table":
            rows = element.find_all("tr")
            if not rows:
                continue
            column_count = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
            table = doc.add_table(rows=0, cols=column_count)
            table.autofit = False
            table.style = "Table Grid"
            widths = _table_widths(element, column_count)
            for row_index, source_row in enumerate(rows):
                target_cells = table.add_row().cells
                for column_index, source_cell in enumerate(source_row.find_all(["th", "td"], recursive=False)):
                    cell = target_cells[column_index]
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    _set_cell_margins(cell)
                    _populate_table_cell(cell, source_cell, row_index)
                    if row_index == 0:
                        _shade(cell, NAVY)
                if row_index == 0:
                    _set_repeat_header(table.rows[0])
            _set_table_geometry(table, widths)


def export_docx(
    html_path: Path,
    output_path: Path,
    approved: bool = False,
    page_map: dict[str, int] | None = None,
) -> Path:
    if not approved:
        raise PermissionError("DOCX export requires explicit user approval (--approved).")
    BeautifulSoup, Document = _dependencies()
    soup = BeautifulSoup(html_path.read_text(encoding="utf-8"), "html.parser")
    meta = {tag.get("name", "")[5:]: tag.get("content", "") for tag in soup.find_all("meta") if tag.get("name", "").startswith("apta:")}
    doc = Document()
    _style_document(doc)
    _disable_update_fields(doc)
    asset_dir = Path(__file__).resolve().parent.parent / "assets"
    _add_cover(doc, meta, asset_dir / "logo_aptaworks_wide_white.png", soup, html_path.parent)
    _add_content(
        doc,
        soup,
        html_path.parent,
        meta,
        asset_dir / "logo_aptaworks_wide_navy.png",
        page_map=page_map,
    )
    core = doc.core_properties
    core.title = meta.get("title", html_path.stem)
    core.subject = meta.get("subtitle", "Aptaworks proposal")
    core.author = meta.get("author", "Aptaworks Team")
    core.keywords = "Aptaworks, proposal"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("html", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--approved", action="store_true", help="Confirm that the user explicitly approved DOCX export.")
    args = parser.parse_args()
    output = (args.output or args.html.with_suffix(".docx")).resolve()
    try:
        result = export_docx(args.html.resolve(), output, approved=args.approved)
    except (OSError, PermissionError, RuntimeError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"DOCX: {result}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
